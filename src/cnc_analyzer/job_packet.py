from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Callable

from .models import JobInfo, JobPacketData, ToolListEntry


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:g}"
    return str(value).strip()


def _norm(value: Any) -> str:
    return " ".join(_text(value).upper().replace("\n", " ").replace(":", " ").split())


def _parse_tool_number(value: Any) -> int | None:
    text = _text(value).upper().strip()
    if text.startswith("T"):
        text = text[1:].strip()
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return None


def _extract_from_grid(
    sheet_names: list[str],
    get_value: Callable[[str, int, int], Any],
    max_row: Callable[[str], int],
    max_col: Callable[[str], int],
    source_filename: str,
) -> JobPacketData:
    name_lookup = {name.strip().upper(): name for name in sheet_names}
    info = JobInfo()

    data_sheet = name_lookup.get("DATA ENTRY")
    if data_sheet:
        label_map = {
            "CUSTOMER": "customer",
            "PART NAME": "part_name",
            "PART NUMBER": "part_number",
            "REVISION": "part_revision",
            "PROGRAM # / PALLET 1": "program_number",
            "PROGRAM # PALLET 1": "program_number",
            "MACHINE #": "machine",
            "OPERATION #": "operation",
        }
        for row in range(1, min(max_row(data_sheet), 60) + 1):
            label = _norm(get_value(data_sheet, row, 1))
            if not label:
                continue
            label = label.replace(" :", "").strip()
            for candidate, attr in label_map.items():
                if label == candidate or label.startswith(candidate):
                    value = _text(get_value(data_sheet, row, 2))
                    if value:
                        setattr(info, attr, value)
                    break

    tool_sheet = name_lookup.get("TOOL LIST")
    process_sheet = name_lookup.get("PROCESS SHEET")

    # Fallback to the visible packet headers when DATA ENTRY is blank or formulas
    # in the old template do not currently have a valid cached result.
    if tool_sheet:
        fallbacks = {
            "customer": (2, 4),
            "part_number": (2, 8),
            "operation": (2, 13),
            "machine": (3, 3),
            "program_number": (3, 14),
            "part_revision": (4, 4),
            "part_name": (4, 7),
        }
        for attr, (row, col) in fallbacks.items():
            if not getattr(info, attr):
                value = _text(get_value(tool_sheet, row, col))
                if value and value not in {"#VALUE!", "#REF!", "#N/A"}:
                    setattr(info, attr, value)

    if process_sheet:
        fallbacks = {
            "customer": (2, 4),
            "part_number": (2, 8),
            "operation": (2, 12),
            "machine": (3, 3),
            "program_number": (3, 12),
        }
        for attr, (row, col) in fallbacks.items():
            if not getattr(info, attr):
                value = _text(get_value(process_sheet, row, col))
                if value and value not in {"#VALUE!", "#REF!", "#N/A"}:
                    setattr(info, attr, value)

    tools: list[ToolListEntry] = []
    if tool_sheet:
        header_row = None
        headers: dict[str, int] = {}
        for row in range(1, min(max_row(tool_sheet), 80) + 1):
            row_headers: dict[str, int] = {}
            for col in range(1, min(max_col(tool_sheet), 40) + 1):
                text = _norm(get_value(tool_sheet, row, col))
                if not text:
                    continue
                if text.startswith("TOOL #"):
                    row_headers["tool"] = col
                elif "DESCRIPTION" in text:
                    row_headers["description"] = col
                elif "HOLDER TYPE" in text:
                    row_headers["holder"] = col
                elif "TOOL PROJECTION" in text:
                    row_headers["projection"] = col
                elif "TOOL LENGTH" in text and "VALUE" in text:
                    row_headers["h_measure"] = col
                elif "CUTTER" in text and "(D)" in text:
                    row_headers["d_value"] = col
                elif "SPECIAL INSTRUCTIONS" in text:
                    row_headers["special"] = col
            if "tool" in row_headers and "description" in row_headers:
                header_row = row
                headers = row_headers
                break

        if header_row is not None:
            blank_run = 0
            for row in range(header_row + 1, min(max_row(tool_sheet), header_row + 150) + 1):
                tool_number = _parse_tool_number(get_value(tool_sheet, row, headers["tool"]))
                if tool_number is None:
                    blank_run += 1
                    if blank_run >= 8:
                        break
                    continue
                blank_run = 0
                tools.append(ToolListEntry(
                    tool_number=tool_number,
                    description=_text(get_value(tool_sheet, row, headers["description"])),
                    holder_type=_text(get_value(tool_sheet, row, headers.get("holder", 0))) if headers.get("holder") else "",
                    projection=_text(get_value(tool_sheet, row, headers.get("projection", 0))) if headers.get("projection") else "",
                    h_measure=_text(get_value(tool_sheet, row, headers.get("h_measure", 0))) if headers.get("h_measure") else "",
                    d_value=_text(get_value(tool_sheet, row, headers.get("d_value", 0))) if headers.get("d_value") else "",
                    special_instructions=_text(get_value(tool_sheet, row, headers.get("special", 0))) if headers.get("special") else "",
                    source="Job Packet",
                ))

    return JobPacketData(job_info=info, tools=tools, source_filename=source_filename)


def _read_openxml(data: bytes, filename: str) -> JobPacketData:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        sheet_names = list(wb.sheetnames)

        def get_value(sheet: str, row: int, col: int) -> Any:
            if row <= 0 or col <= 0:
                return None
            return wb[sheet].cell(row=row, column=col).value

        def max_row(sheet: str) -> int:
            return wb[sheet].max_row

        def max_col(sheet: str) -> int:
            return wb[sheet].max_column

        return _extract_from_grid(sheet_names, get_value, max_row, max_col, filename)
    finally:
        wb.close()


def _read_xls(data: bytes, filename: str) -> JobPacketData:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - installed by requirements in the packaged app
        raise RuntimeError(
            "Reading legacy .xls job packets requires the xlrd package. Re-run the launcher so dependencies are installed."
        ) from exc

    wb = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        sheet_names = wb.sheet_names()

        def get_value(sheet: str, row: int, col: int) -> Any:
            if row <= 0 or col <= 0:
                return None
            ws = wb.sheet_by_name(sheet)
            if row > ws.nrows or col > ws.ncols:
                return None
            return ws.cell_value(row - 1, col - 1)

        def max_row(sheet: str) -> int:
            return wb.sheet_by_name(sheet).nrows

        def max_col(sheet: str) -> int:
            return wb.sheet_by_name(sheet).ncols

        return _extract_from_grid(sheet_names, get_value, max_row, max_col, filename)
    finally:
        wb.release_resources()


def read_job_packet_bytes(data: bytes, filename: str) -> JobPacketData:
    if not data:
        raise ValueError("The selected job-packet file is empty.")

    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_openxml(data, filename)
    if suffix == ".xls":
        return _read_xls(data, filename)
    raise ValueError("Job packet must be an Excel .xls, .xlsx, or .xlsm file.")
