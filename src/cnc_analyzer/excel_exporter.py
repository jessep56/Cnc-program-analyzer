from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import AnalysisResult, JobInfo, ToolListEntry, WarningItem
from .reporting import process_sheet_rows, technical_sequence_rows, tool_rows, warning_rows

NAVY = "17324D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_RED = "FCE8E6"
LIGHT_YELLOW = "FFF4CC"
LIGHT_GRAY = "F3F5F7"
LIGHT_GREEN = "E6F4EA"
WHITE = "FFFFFF"
DARK = "1F2937"
BORDER = Side(style="thin", color="B7C0C8")

PROCESS_WIDTHS = {
    "Step": 6,
    "Tool": 7,
    "Tool Description": 27,
    "Process / Feature": 29,
    "Qty": 7,
    "WCS": 12,
    "Offsets / Comp": 18,
    "Lines": 10,
    "Notes": 18,
}

TOOL_LIST_END_ROW = 106

TOOL_LIST_HEADERS = [
    "Tool #", "Description", "Holder Type", "Projection", "H Measure", "D Value", "Special Instructions"
]

CENTER_HEADERS = {
    "Step", "Tool", "Qty", "WCS", "Offsets / Comp", "Lines",
    "Tool Changes", "Sequences", "Cycle Locations", "H Offsets", "D Offsets",
}


def _safe(value: str | None) -> str:
    return (value or "").strip()


def _effective_job_info(result: AnalysisResult, job_info: JobInfo | None) -> JobInfo:
    job_info = job_info or JobInfo()
    return JobInfo(
        customer=_safe(job_info.customer),
        part_number=_safe(job_info.part_number),
        part_name=_safe(job_info.part_name),
        part_revision=_safe(job_info.part_revision),
        program_number=_safe(job_info.program_number) or result.program_name,
        machine=_safe(job_info.machine),
        operation=_safe(job_info.operation),
        automation_notes=_safe(job_info.automation_notes),
    )


def _configure_page(ws, *, zoom: int = 100, fit_wide: bool = False, freeze: str = "A2") -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.sheet_view.zoomScaleNormal = zoom
    ws.freeze_panes = freeze
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    if fit_wide:
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0


def _label_style(cell) -> None:
    cell.font = Font(bold=True, color=DARK, size=10)
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)


def _value_style(cell) -> None:
    cell.font = Font(color=DARK, size=10)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)


def _add_logo(ws) -> None:
    logo_path = Path(__file__).resolve().parents[2] / "assets" / "cwm_logo.png"
    if not logo_path.exists():
        return
    image = Image(str(logo_path))
    image.width = 94
    image.height = 44
    ws.add_image(image, "A1")


def _write_job_header(ws, title: str, info: JobInfo, *, end_col: int = 9) -> None:
    end_letter = get_column_letter(end_col)
    _add_logo(ws)

    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=end_col)
    title_cell = ws.cell(row=1, column=3, value=title)
    title_cell.font = Font(size=18, bold=True, color=WHITE)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 31
    for row in (2, 3, 4):
        ws.row_dimensions[row].height = 21

    # Logo occupies A1:B4. Job information stays compact on the right.
    pairs = [
        (2, 3, "Customer:", 4, 5, info.customer),
        (2, 6, "Part Number:", 7, 9, info.part_number),
        (3, 3, "Part Name:", 4, 5, info.part_name),
        (3, 6, "Part Rev:", 7, 7, info.part_revision),
        (3, 8, "Program:", 9, 9, info.program_number),
        (4, 3, "Machine:", 4, 4, info.machine),
        (4, 5, "Operation:", 6, 6, info.operation),
        (4, 7, "Analyzed:", 8, 9, datetime.now().strftime("%Y-%m-%d")),
    ]
    for row, label_col, label, value_start, value_end, value in pairs:
        label_cell = ws.cell(row=row, column=label_col, value=label)
        _label_style(label_cell)
        if value_end > value_start:
            ws.merge_cells(start_row=row, start_column=value_start, end_row=row, end_column=value_end)
        value_cell = ws.cell(row=row, column=value_start, value=value)
        _value_style(value_cell)
        for col in range(value_start + 1, value_end + 1):
            _value_style(ws.cell(row=row, column=col))

    # Border/fill the logo footprint so the header still looks deliberate if the
    # image is unavailable on a future platform.
    for row in range(1, 5):
        for col in range(1, 3):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
            cell.fill = PatternFill("solid", fgColor=WHITE)

    # Preserve the title fill across merged cells.
    for col in range(3, end_col + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=NAVY)


def _add_table(ws, ref: str, name: str, style: str = "TableStyleMedium2") -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=style,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _write_process_sheet(ws, result: AnalysisResult, info: JobInfo, tools: list[ToolListEntry]) -> None:
    _configure_page(ws, zoom=110, fit_wide=True, freeze="A7")
    _write_job_header(ws, "CNC PROGRAM PROCESS SHEET", info, end_col=9)
    ws.print_title_rows = "1:6"

    rows = process_sheet_rows(result, tools)
    headers = list(rows[0].keys()) if rows else [
        "Step", "Tool", "Tool Description", "Process / Feature", "Qty", "WCS", "Offsets / Comp", "Lines", "Notes"
    ]
    header_row = 6
    data_start = 7

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
    ws.row_dimensions[header_row].height = 27

    for row_index, row in enumerate(rows, start=data_start):
        for col_index, header in enumerate(headers, start=1):
            if header == "Tool Description":
                # Formula-driven by design: if setup changes T7 on the Tool List later,
                # the Process Sheet updates without regenerating the analyzer report.
                tool_expr = f'VALUE(SUBSTITUTE(B{row_index},"T",""))'
                formula = (
                    f'=IF(COUNTIF(\'Tool List\'!$A$7:$A${TOOL_LIST_END_ROW},{tool_expr})=0,'
                    '"⚠ TOOL NOT FOUND",'
                    f'INDEX(\'Tool List\'!$B$7:$B${TOOL_LIST_END_ROW},MATCH({tool_expr},\'Tool List\'!$A$7:$A${TOOL_LIST_END_ROW},0)))'
                )
                cell = ws.cell(row=row_index, column=col_index, value=formula)
            else:
                cell = ws.cell(row=row_index, column=col_index, value=row.get(header, ""))
            horizontal = "center" if header in CENTER_HEADERS else "left"
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
        ws.row_dimensions[row_index].height = 25

    end_row = max(header_row + 1, data_start + len(rows) - 1)
    end_col = get_column_letter(len(headers))
    if rows:
        _add_table(ws, f"A{header_row}:{end_col}{end_row}", "ProcessSheetTable")

    for col_index, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = PROCESS_WIDTHS.get(header, 14)

    if rows:
        desc_range = f"C{data_start}:C{end_row}"
        ws.conditional_formatting.add(
            desc_range,
            FormulaRule(
                formula=[f'C{data_start}="⚠ TOOL NOT FOUND"'],
                fill=PatternFill("solid", fgColor=LIGHT_RED),
                font=Font(color="9C0006", bold=True),
            ),
        )

    last_content_row = end_row
    if info.automation_notes:
        note_title_row = end_row + 2
        note_row = note_title_row + 1
        ws.merge_cells(start_row=note_title_row, start_column=1, end_row=note_title_row, end_column=9)
        ws.cell(row=note_title_row, column=1, value="AUTOMATION / SPECIAL INSTRUCTIONS")
        ws.cell(row=note_title_row, column=1).font = Font(bold=True, color=WHITE)
        ws.cell(row=note_title_row, column=1).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row=note_title_row, column=1).alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=9)
        ws.cell(row=note_row, column=1, value=info.automation_notes)
        ws.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=note_row, column=1).fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
        ws.row_dimensions[note_row].height = 35
        last_content_row = note_row + 2

    ws.print_area = f"A1:I{last_content_row}"
    ws.oddFooter.center.text = "Verification aid only — verify tooling, offsets, workholding, and machine state before production."
    ws.oddFooter.center.size = 8


def _write_tool_list(ws, info: JobInfo, tools: list[ToolListEntry]) -> None:
    _configure_page(ws, zoom=105, fit_wide=True, freeze="A7")
    _write_job_header(ws, "TOOL LIST", info, end_col=7)
    ws.print_title_rows = "1:6"

    header_row = 6
    for col, header in enumerate(TOOL_LIST_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)

    sorted_tools = sorted(tools, key=lambda tool: tool.tool_number)
    last_tool_row = 6
    for row_offset, tool in enumerate(sorted_tools):
        row = 7 + row_offset
        if row > TOOL_LIST_END_ROW:
            break
        values = [
            tool.tool_number,
            tool.description,
            tool.holder_type,
            tool.projection,
            tool.h_measure,
            tool.d_value,
            tool.special_instructions,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(
                horizontal="center" if col in {1, 4, 5, 6} else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
        ws.row_dimensions[row].height = 22
        last_tool_row = row

    # Keep the lookup range larger than the visible table so users can add a new tool
    # later without regenerating the workbook.
    for row in range(last_tool_row + 1, TOOL_LIST_END_ROW + 1):
        for col in range(1, 8):
            ws.cell(row=row, column=col).alignment = Alignment(
                horizontal="center" if col in {1, 4, 5, 6} else "left",
                vertical="center",
                wrap_text=True,
            )

    if sorted_tools:
        _add_table(ws, f"A6:G{last_tool_row}", "ToolListTable", style="TableStyleMedium2")

    widths = [8, 35, 18, 14, 14, 15, 28]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # A populated tool number with a blank description should stand out immediately.
    ws.conditional_formatting.add(
        f"B7:B{TOOL_LIST_END_ROW}",
        FormulaRule(
            formula=['AND($A7<>"",$B7="")'],
            fill=PatternFill("solid", fgColor=LIGHT_YELLOW),
            font=Font(color="9C6500", bold=True),
        ),
    )
    ws.print_area = f"A1:G{max(last_tool_row, 7)}"


def _write_table_sheet(
    ws,
    rows: list[dict[str, object]],
    table_name: str,
    *,
    zoom: int = 100,
    fit_wide: bool = False,
) -> None:
    _configure_page(ws, zoom=zoom, fit_wide=fit_wide)
    if not rows:
        ws["A1"] = "No records"
        return

    headers = list(rows[0].keys())
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=BORDER)
    ws.row_dimensions[1].height = 24

    for row_index, row in enumerate(rows, start=2):
        for col_index, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=row.get(header, ""))
            horizontal = "center" if header in CENTER_HEADERS else "left"
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
            cell.border = Border(bottom=BORDER)
        ws.row_dimensions[row_index].height = 24

    end_col = get_column_letter(len(headers))
    end_row = len(rows) + 1
    _add_table(ws, f"A1:{end_col}{end_row}", table_name)

    for col_index, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row_index in range(2, end_row + 1):
            max_len = max(max_len, len(str(ws.cell(row=row_index, column=col_index).value or "")))
        width = min(max(max_len + 2, 9), 30)
        if header in {"Description", "Message", "Notes", "Operations", "Features", "Program Tool Description", "Feature"}:
            width = min(max(width, 22), 38)
        ws.column_dimensions[get_column_letter(col_index)].width = width

    if fit_wide:
        ws.print_area = f"A1:{end_col}{end_row}"


def _write_summary(ws, result: AnalysisResult, info: JobInfo, warnings: list[WarningItem]) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 105
    ws.merge_cells("A1:F1")
    ws["A1"] = "CNC Program Analyzer — Analysis Summary"
    ws["A1"].font = Font(size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    summary_data = [
        ("Customer", info.customer),
        ("Part Number", info.part_number),
        ("Part Name", info.part_name),
        ("Revision", info.part_revision),
        ("Program", info.program_number or result.program_name),
        ("Machine", info.machine),
        ("Source file", result.filename),
        ("Controller model", result.controller_family),
        ("Source lines", result.line_count),
        ("Tool changes", len(result.tool_changes)),
        ("Unique tools", result.unique_tools),
        ("Machining sequences", len(result.sequences)),
        ("Canned-cycle locations", result.canned_cycle_locations),
        ("Warnings / notices", len(warnings)),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]

    row = 3
    for label, value in summary_data:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color=DARK)
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row=row, column=2, value=value)
        row += 1

    ws.cell(row=row + 1, column=1, value="Important")
    ws.cell(row=row + 1, column=1).font = Font(bold=True, color=WHITE)
    ws.cell(row=row + 1, column=1).fill = PatternFill("solid", fgColor=BLUE)
    ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 4, end_column=6)
    ws.cell(row=row + 2, column=1, value=(
        "This report is a verification aid, not a substitute for machine prove-out, setup verification, "
        "or review by qualified personnel. Macro/subprogram behavior may require manual review."
    ))
    ws.cell(row=row + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=row + 2, column=1).fill = PatternFill("solid", fgColor=LIGHT_YELLOW)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 36
    for col in "CDEF":
        ws.column_dimensions[col].width = 14


def build_excel_report(
    result: AnalysisResult,
    *,
    job_info: JobInfo | None = None,
    tools: list[ToolListEntry] | None = None,
    warnings: list[WarningItem] | None = None,
) -> bytes:
    info = _effective_job_info(result, job_info)
    tools = tools or []
    warnings = warnings if warnings is not None else result.warnings

    wb = Workbook()
    process_ws = wb.active
    process_ws.title = "Process Sheet"
    _write_process_sheet(process_ws, result, info, tools)

    tool_list_ws = wb.create_sheet("Tool List")
    _write_tool_list(tool_list_ws, info, tools)

    warnings_ws = wb.create_sheet("Warnings")
    warning_data = warning_rows(warnings)
    _write_table_sheet(warnings_ws, warning_data, "WarningsTable", zoom=100)
    if warning_data:
        for row_index in range(2, len(warning_data) + 2):
            severity = warnings_ws.cell(row=row_index, column=1).value
            fill = None
            if severity == "Error":
                fill = LIGHT_RED
            elif severity == "Warning":
                fill = LIGHT_YELLOW
            elif severity == "Info":
                fill = LIGHT_GRAY
            if fill:
                for col_index in range(1, 6):
                    warnings_ws.cell(row=row_index, column=col_index).fill = PatternFill("solid", fgColor=fill)

    summary_ws = wb.create_sheet("Summary")
    _write_summary(summary_ws, result, info, warnings)

    technical_ws = wb.create_sheet("Technical Details")
    _write_table_sheet(technical_ws, technical_sequence_rows(result), "TechnicalDetailsTable", zoom=90)

    tool_summary_ws = wb.create_sheet("Tool Summary")
    _write_table_sheet(tool_summary_ws, tool_rows(result, tools), "ToolSummaryTable", zoom=100)

    # Force Excel/iPad Excel to refresh Tool Description formulas on open.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
